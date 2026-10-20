"""Turn an uploaded CV file (PDF, DOCX or TXT) or pasted text into clean plain text."""
from io import BytesIO

from docx import Document
from fastapi import HTTPException, UploadFile
from pypdf import PdfReader

ALLOWED_TYPES = (".pdf", ".docx", ".txt", ".md")
MAX_FILE_SIZE = 4 * 1024 * 1024   # 4 MB (Vercel rejects requests above ~4.5 MB)
MAX_CV_CHARS = 20_000             # enough for a long CV, protects token costs
MIN_CV_CHARS = 100                # less than this is not a real CV


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()

    if name.endswith(".pdf"):
        reader = PdfReader(BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if name.endswith(".docx"):
        doc = Document(BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        # Many CVs use tables for layout, so read those too
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)

    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="ignore")

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


def get_cv_text(file: UploadFile | None, cv_text: str) -> str:
    """
    Shared helper for every feature that needs a CV.
    Uses the uploaded file if there is one, otherwise the pasted text.
    Raises HTTPException(400) with a clear message if something is wrong.
    """
    # 1. A file was uploaded: check it and read the text out of it
    if file is not None and file.filename:
        if not file.filename.lower().endswith(ALLOWED_TYPES):
            raise HTTPException(400, "Please upload a PDF, DOCX or TXT file.")

        data = file.file.read()
        if len(data) > MAX_FILE_SIZE:
            raise HTTPException(400, "File is too big. The limit is 4 MB.")

        try:
            cv_text = extract_text(file.filename, data)
        except Exception:
            raise HTTPException(400, "Could not read this file. Try another format or paste the text.")

    cv_text = (cv_text or "").strip()

    # 2. Check we really have a CV
    if len(cv_text) < MIN_CV_CHARS:
        raise HTTPException(
            400,
            "Not enough text found. If your PDF is a scanned image, paste the CV as text instead.",
        )

    # 3. Cut very long CVs so the AI call stays cheap
    return cv_text[:MAX_CV_CHARS]
